#! /usr/bin/env python
# @Time    : 2026/9/28
# @Author  : afish
# @File    : test_header_footer.py
"""页眉页脚检查/修正（header_footer.py）单元测试（issue#93）。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from wordformat.header_footer import (
    DEFAULT_HF_CONFIG,
    apply_header_footer,
    check_header_footer,
    resolve_hf_config,
)


def _set_east_asia(run, font_name: str) -> None:
    """手动设置 run 的中文字体（python-docx 不支持 eastAsia）。"""
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)


def _header_text(doc, text, size_pt=9.0, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                 en_font="Times New Roman", cn_font="宋体"):
    """构造带指定格式的页眉段落并返回其 run。"""
    section = doc.sections[0]
    para = section.header.paragraphs[0]
    run = para.add_run(text)
    run.font.size = Pt(size_pt)
    run.font.name = en_font
    _set_east_asia(run, cn_font)
    para.alignment = alignment
    return run


class TestResolveHfConfig:
    def test_none_returns_default(self):
        cfg = resolve_hf_config(None)
        assert cfg == DEFAULT_HF_CONFIG

    def test_dict_merges_user_config(self):
        cfg = resolve_hf_config({"header_footer": {"header": {"font_size": "五号"}}})
        assert cfg["header"]["font_size"] == "五号"
        assert cfg["footer"]["font_size"] == "小五"  # 深合并保留缺省

    def test_disabled_returns_empty(self):
        assert resolve_hf_config({"header_footer": {"enabled": False}}) == {}

    def test_config_object_attribute(self):
        class FakeConfig:
            header_footer = {"footer": {"chinese_font_name": "黑体"}}

        cfg = resolve_hf_config(FakeConfig())
        assert cfg["footer"]["chinese_font_name"] == "黑体"


class TestCheckHeaderFooter:
    def test_ok_header_no_issues(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文")
        assert check_header_footer(doc) == []

    def test_wrong_alignment_reported(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文", alignment=WD_ALIGN_PARAGRAPH.LEFT)
        issues = check_header_footer(doc)
        assert any("对齐错误" in t and "页眉(第1节)" in t for t in issues)

    def test_wrong_font_size_reported(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文", size_pt=10.5)  # 五号 ≠ 小五
        issues = check_header_footer(doc)
        assert any("字号错误" in t for t in issues)

    def test_wrong_fonts_reported(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文", en_font="Arial", cn_font="黑体")
        issues = check_header_footer(doc)
        texts = "\n".join(issues)
        assert "中文字体错误" in texts
        assert "英文字体错误" in texts

    def test_empty_header_with_require_content(self):
        cfg = {"header_footer": {"header": {"require_content": True}}}
        doc = Document()  # 页眉为空
        issues = check_header_footer(doc, cfg)
        assert any("内容缺失" in t for t in issues)

    def test_empty_header_without_require_content(self):
        doc = Document()
        assert check_header_footer(doc) == []

    def test_linked_to_previous_section_skipped(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文", size_pt=10.5)
        doc.add_section()
        # 第二节沿用第一节（is_linked_to_previous=True），其页眉不参与检查
        assert doc.sections[1].header.is_linked_to_previous is True
        issues = check_header_footer(doc)
        # 只有第一节问题；第二节标题不带 (第2节) 前缀
        assert len(issues) == 1

    def test_disabled_config_skips(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文", size_pt=10.5)
        assert check_header_footer(doc, {"header_footer": {"enabled": False}}) == []


class TestApplyHeaderFooter:
    def test_apply_fixes_format_no_remaining(self):
        doc = Document()
        _header_text(
            doc,
            "某某大学毕业论文",
            size_pt=10.5,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            en_font="Arial",
            cn_font="黑体",
        )
        remaining = apply_header_footer(doc)
        assert remaining == []
        assert check_header_footer(doc) == []

    def test_apply_empty_header_leaves_remaining_empty(self):
        doc = Document()  # 页眉为空：不自动补内容，无残留
        assert apply_header_footer(doc) == []

    def test_apply_disabled_config_noop(self):
        doc = Document()
        run = _header_text(doc, "某某大学毕业论文", size_pt=10.5)
        assert apply_header_footer(doc, {"header_footer": {"enabled": False}}) == []
        assert run.font.size.pt == 10.5

    def test_apply_second_section_linked_skipped(self):
        doc = Document()
        _header_text(doc, "某某大学毕业论文", size_pt=10.5)
        doc.add_section()
        apply_header_footer(doc)
        # 第二节沿用第一节（未指认新内容），修正后第一节字号为 9pt
        assert doc.sections[0].header.paragraphs[0].runs[0].font.size.pt == 9.0
